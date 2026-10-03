#!/usr/bin/env python3
"""Routes Flask — templates, IMAP/filtres, marque-réponse, exports/imports (blueprint)."""
import io, csv, json, os, re, socket, time, threading, uuid
from datetime import datetime, timedelta
from pathlib import Path

from flask import (Blueprint, jsonify, request, render_template, Response,
                   send_file, session, redirect)

from candimail.core import db, crypto
from candimail.core.config import (BASE_DIR, UPLOADS_DIR, CONFIG_FILE, LOG_FILE,
                        DEFAULT_CONFIG, DEFAULT_CV, TEMPLATES)
from candimail.core.noyau import (current_uid, charger_config, sauver_config, _config_globale,
                   charger_log, sauver_log, get_delai, check_rate_limit,
                   maintenant_str, parse_emails, get_comptes_actifs,
                   compte_suivant, _get_template, build_letter, build_subject,
                   get_pieces_actives, recharger_pieces, _pieces,
                   _config_lock, _config_cache, _log_lock, _send_lock,
                   _log_send_lock, _bulk_lock, _bulk_running, _bulk_state,
                   _compte_index, _auth_attempts, _send_timestamps,
                   _last_send_time, _txt, _email_valide)
from candimail.services.smtp import envoyer_email, _envoyer_un_email
from candimail.services.imap import (check_imap_compte, verifier_toutes_boites,
                           appliquer_filtres_compte, _filtre_valide,
                           _ACTIONS_VALIDES, _creer_libelle, _utf7)

bp = Blueprint("divers", __name__)


# ------------------------------------------------------------
# TEMPLATES PERSONNALISÉS
# ------------------------------------------------------------
@bp.route("/api/templates", methods=["GET"])
def api_templates_list():
    cfg = charger_config()
    return jsonify({"templates": cfg.get("templates_perso", [])})

@bp.route("/api/templates", methods=["POST"])
def api_templates_save():
    data = request.get_json() or {}
    nom       = (data.get("nom") or "").strip()
    objet     = (data.get("objet") or "").strip()
    corps     = (data.get("corps") or "").strip()
    categorie = (data.get("categorie") or "autre").strip()
    tpl_id    = (data.get("id") or "").strip()

    if not nom or not objet or not corps:
        return jsonify({"ok": False, "error": "Nom, objet et corps requis"}), 400

    cfg = charger_config()
    templates = cfg.get("templates_perso", [])

    if tpl_id:
        found = False
        for t in templates:
            if t["id"] == tpl_id:
                t["nom"]       = nom
                t["objet"]     = objet
                t["corps"]     = corps
                t["categorie"] = categorie
                t["maj_le"]    = maintenant_str()
                found = True
                break
        if not found:
            return jsonify({"ok": False, "error": "Template introuvable"}), 404
    else:
        new_id = f"tpl_{int(time.time())}_{uuid.uuid4().hex[:6]}"
        templates.append({
            "id": new_id, "nom": nom, "objet": objet,
            "corps": corps, "categorie": categorie,
            "cree_le": maintenant_str(),
        })
        tpl_id = new_id

    cfg["templates_perso"] = templates
    sauver_config(cfg)
    return jsonify({"ok": True, "id": tpl_id})

@bp.route("/api/templates/<tpl_id>", methods=["DELETE"])
def api_templates_delete(tpl_id):
    cfg = charger_config()
    templates = cfg.get("templates_perso", [])
    avant = len(templates)
    templates = [t for t in templates if t["id"] != tpl_id]
    if len(templates) == avant:
        return jsonify({"ok": False, "error": "Introuvable"}), 404
    cfg["templates_perso"] = templates
    sauver_config(cfg)
    return jsonify({"ok": True})

@bp.route("/api/templates/<tpl_id>/duplicate", methods=["POST"])
def api_templates_duplicate(tpl_id):
    cfg = charger_config()
    templates = cfg.get("templates_perso", [])
    src = next((t for t in templates if t["id"] == tpl_id), None)
    if not src:
        return jsonify({"ok": False, "error": "Introuvable"}), 404
    new_id = f"tpl_{int(time.time())}_{uuid.uuid4().hex[:6]}"
    copie = {
        "id": new_id, "nom": f"{src['nom']} (copie)",
        "objet": src["objet"], "corps": src["corps"],
        "categorie": src.get("categorie", "autre"),
        "cree_le": maintenant_str(),
    }
    templates.append(copie)
    cfg["templates_perso"] = templates
    sauver_config(cfg)
    return jsonify({"ok": True, "id": new_id})

# ------------------------------------------------------------
# IMAP
# ------------------------------------------------------------
@bp.route("/api/imap/check", methods=["POST"])
def api_imap_check():
    data = request.get_json() or {}
    try:
        jours = min(max(int(data.get("jours", 30)), 1), 365)
    except (TypeError, ValueError):
        jours = 30
    res = verifier_toutes_boites(jours=jours, uid=current_uid())
    return jsonify(res)

@bp.route("/api/imap/status")
def api_imap_status():
    cfg = charger_config()
    return jsonify({
        "derniere_verif": cfg["imap"].get("derniere_verif", 0),
        "auto_verif":     cfg["imap"].get("auto_verif", True),
        "intervalle_min": cfg["imap"].get("intervalle_min", 30),
    })

# Règles de classement par défaut (priorité 1 = évaluée en premier)
# Mots-clés réels : expéditeurs + sujets courants, français et anglais.
REGLES_DEFAUT = [
    {"name": "Bourse",       "icon": "📈", "color": "#f39c12",
     "items": [{"champ": "from",   "valeur": "boursorama"},
                {"champ": "from",   "valeur": "saxobank"},
                {"champ": "from",   "valeur": "trading212"},
                {"champ": "from",   "valeur": "degiro"},
                {"champ": "from",   "valeur": "e*trade"},
                {"champ": "from",   "valeur": "interactivebrokers"},
                {"champ": "sujet",  "valeur": "bourse"},
                {"champ": "sujet",  "valeur": "trading"},
                {"champ": "sujet",  "valeur": "investir"},
                {"champ": "sujet",  "valeur": "investissement"},
                {"champ": "sujet",  "valeur": "portefeuille"},
                {"champ": "sujet",  "valeur": "actions"},
                {"champ": "sujet",  "valeur": "dividende"},
                {"champ": "sujet",  "valeur": "cotations"}],
     "notify": False},
    {"name": "Opportunité",  "icon": "💼", "color": "#27ae60",
     "items": [{"champ": "sujet",  "valeur": "stage"},
                {"champ": "sujet",  "valeur": "stagiaire"},
                {"champ": "sujet",  "valeur": "alternance"},
                {"champ": "sujet",  "valeur": "apprenti"},
                {"champ": "sujet",  "valeur": "opportunité"},
                {"champ": "sujet",  "valeur": "candidature"},
                {"champ": "sujet",  "valeur": "poste"},
                {"champ": "sujet",  "valeur": "recrutement"},
                {"champ": "sujet",  "valeur": "internship"},
                {"champ": "sujet",  "valeur": "opportunity"},
                {"champ": "sujet",  "valeur": "hiring"},
                {"champ": "from",   "valeur": "recrutement"},
                {"champ": "from",   "valeur": "recrute"},
                {"champ": "from",   "valeur": "rh@"},
                {"champ": "from",   "valeur": "jobs@"},
                {"champ": "from",   "valeur": "career"},
                {"champ": "from",   "valeur": "talent"}],
     "notify": True},
    {"name": "Travail",      "icon": "🏢", "color": "#1E5AA8",
     "items": [{"champ": "sujet",  "valeur": "entretien"},
                {"champ": "sujet",  "valeur": "contrat"},
                {"champ": "sujet",  "valeur": "embauche"},
                {"champ": "sujet",  "valeur": "offre d'emploi"},
                {"champ": "sujet",  "valeur": "interview"},
                {"champ": "sujet",  "valeur": "job offer"},
                {"champ": "sujet",  "valeur": "convention de stage"},
                {"champ": "sujet",  "valeur": "votre candidature"},
                {"champ": "from",   "valeur": "noreply@linkedin.com"},
                {"champ": "from",   "valeur": "welcome@linkedin.com"},
                {"champ": "from",   "valeur": "notifications@linkedin.com"},
                {"champ": "from",   "valeur": "hello@indeed.com"},
                {"champ": "from",   "valeur": "welcomemail@welcomemail.welcomekit.co"},
                {"champ": "from",   "valeur": "welcomekit"},
                {"champ": "from",   "valeur": "jobteaser"},
                {"champ": "from",   "valeur": "jobvite"},
                {"champ": "from",   "valeur": "greenhouse.io"},
                {"champ": "from",   "valeur": "workday"},
                {"champ": "from",   "valeur": "lever.co"},
                {"champ": "from",   "valeur": "smartrecruiters"}],
     "notify": True},
    {"name": "Université",   "icon": "🎓", "color": "#8e44ad",
     "items": [{"champ": "from",   "valeur": "univ-"},
                {"champ": "from",   "valeur": "universite"},
                {"champ": "from",   "valeur": "university"},
                {"champ": "from",   "valeur": ".edu"},
                {"champ": "from",   "valeur": "esisalat"},
                {"champ": "from",   "valeur": "imesa"},
                {"champ": "sujet",  "valeur": "examen"},
                {"champ": "sujet",  "valeur": "cours"},
                {"champ": "sujet",  "valeur": "étudiant"},
                {"champ": "sujet",  "valeur": "scolarité"},
                {"champ": "sujet",  "valeur": "mémoire"},
                {"champ": "sujet",  "valeur": "soutenance"},
                {"champ": "sujet",  "valeur": "moodle"},
                {"champ": "sujet",  "valeur": "exam"}],
     "notify": False},
    {"name": "Réseaux sociaux", "icon": "👥", "color": "#2980b9",
     "items": [{"champ": "from",   "valeur": "linkedin.com"},
                {"champ": "from",   "valeur": "facebookmail.com"},
                {"champ": "from",   "valeur": "x.com"},
                {"champ": "from",   "valeur": "twitter"},
                {"champ": "from",   "valeur": "instagram"},
                {"champ": "from",   "valeur": "tiktok"},
                {"champ": "from",   "valeur": "whatsapp"},
                {"champ": "from",   "valeur": "telegram"},
                {"champ": "from",   "valeur": "snapchat"}],
     "notify": False},
    {"name": "Annonce",      "icon": "📢", "color": "#e67e22",
     "items": [{"champ": "sujet",  "valeur": "newsletter"},
                {"champ": "sujet",  "valeur": "promo"},
                {"champ": "sujet",  "valeur": "solde"},
                {"champ": "sujet",  "valeur": "réduction"},
                {"champ": "sujet",  "valeur": "vente flash"},
                {"champ": "sujet",  "valeur": "offres exclusives"},
                {"champ": "from",   "valeur": "newsletter"},
                {"champ": "from",   "valeur": "promo"},
                {"champ": "from",   "valeur": "marketing"},
                {"champ": "from",   "valeur": "noreply"}],
     "notify": False},
]


def _seed_regles(cfg):
    """Crée les 6 règles par défaut au premier lancement.
    Version 2 : met à jour les mots-clés des règles par défaut existantes
    (sans toucher aux règles perso ni aux règles d'apprentissage)."""
    if not cfg.get("filtres_init"):
        regles = []
        for i, d in enumerate(REGLES_DEFAUT, 1):
            regles.append({
                "id":       f"rule-{i:03d}",
                "priority": i,
                "enabled":  True,
                "name":     d["name"],
                "icon":     d["icon"],
                "color":    d["color"],
                "items":    d["items"],
                "label":    d["name"],
                "actions":  ["label"] + (["notify"] if d["notify"] else []),
                "notify":   d["notify"],
            })
        cfg["filtres_gmail"] = regles
        cfg["filtres_init"] = 2
        return
    if cfg.get("filtres_init", 0) < 2:
        noms_defaut = {d["name"] for d in REGLES_DEFAUT}
        for d in REGLES_DEFAUT:
            for f in cfg.get("filtres_gmail", []):
                if f.get("name") == d["name"] and f.get("source") != "apprentissage":
                    f["items"] = d["items"]
                    f["icon"] = d["icon"]
                    f["color"] = d["color"]
                    f["notify"] = d["notify"]
                    acts = [a for a in (f.get("actions") or []) if a != "notify"]
                    if d["notify"]:
                        acts.append("notify")
                    f["actions"] = acts
        cfg["filtres_init"] = 2


def _normaliser_regle(data, existante=None):
    """Construit/normalise une règle depuis un payload front."""
    r = dict(existante or {})
    r.setdefault("id", f"rule-{uuid.uuid4().hex[:8]}")
    items = data.get("items")
    if isinstance(items, list) and items:
        propres = []
        for it in items:
            if not isinstance(it, dict):
                continue
            champ = it.get("champ") if it.get("champ") in ("from", "sujet", "motcle") else None
            val = (it.get("valeur") or "").strip()
            if champ and val:
                propres.append({"champ": champ, "valeur": val})
        r["items"] = propres
    else:
        champ = _txt(data.get("champ")) or "motcle"
        val = _txt(data.get("valeur")).strip()
        r["items"] = ([{"champ": champ, "valeur": val}] if val
                      else r.get("items", []))
    r["name"]  = _txt(data.get("name")).strip() or r.get("name", "règle")
    r["label"] = _txt(data.get("label")).strip() or r.get("label", "")
    if data.get("icon"):
        r["icon"] = _txt(data.get("icon"))
    if data.get("color"):
        r["color"] = _txt(data.get("color"))
    if "enabled" in data:
        r["enabled"] = bool(data.get("enabled"))
    else:
        r.setdefault("enabled", True)
    if isinstance(data.get("priority"), int):
        r["priority"] = data["priority"]
    else:
        r.setdefault("priority", 999)
    acts = data.get("actions") if isinstance(data.get("actions"), list) else r.get("actions", [])
    acts = [a for a in acts if a in _ACTIONS_VALIDES]
    r["actions"] = acts
    r["notify"] = "notify" in acts
    return r


@bp.route("/api/filtres", methods=["GET"])
def api_filtres_get():
    cfg = charger_config()
    _seed_regles(cfg)
    if "filtres_init" in cfg:
        sauver_config(cfg)
    return jsonify({"ok": True, "filtres": cfg.get("filtres_gmail", [])})


@bp.route("/api/filtres", methods=["POST"])
def api_filtres_add():
    data = request.get_json() or {}
    cfg = charger_config()
    filtre = _normaliser_regle(data)
    filtre["priority"] = len(cfg.get("filtres_gmail", [])) + 1
    if not _filtre_valide(filtre):
        return jsonify({"ok": False,
                        "error": "Règle invalide (champ, valeur ou actions)"}), 400
    cfg.setdefault("filtres_gmail", []).append(filtre)
    sauver_config(cfg)
    return jsonify({"ok": True, "filtres": cfg["filtres_gmail"]})


@bp.route("/api/filtres/<int:idx>", methods=["PUT", "PATCH"])
def api_filtres_update(idx):
    """Modifie une règle existante (nom, conditions, libellé, actions…).
    La priorité n'est PAS modifiable ici (drag & drop la gère)."""
    data = request.get_json() or {}
    cfg = charger_config()
    filtres = cfg.get("filtres_gmail", [])
    if not 0 <= idx < len(filtres):
        return jsonify({"ok": False, "error": "Introuvable"}), 404
    data.pop("priority", None)
    data.pop("id", None)
    mise_a_jour = _normaliser_regle(data, existante=filtres[idx])
    if not _filtre_valide(mise_a_jour):
        return jsonify({"ok": False,
                        "error": "Règle invalide (champ, valeur ou actions)"}), 400
    filtres[idx] = mise_a_jour
    sauver_config(cfg)
    return jsonify({"ok": True, "filtres": filtres})


@bp.route("/api/filtres/reorder", methods=["POST"])
def api_filtres_reorder():
    """Drag & drop : reçoit la liste des ids dans le nouvel ordre."""
    data = request.get_json() or {}
    ordre = data.get("ordre")
    if not isinstance(ordre, list):
        return jsonify({"ok": False, "error": "ordre requis"}), 400
    cfg = charger_config()
    filtres = cfg.get("filtres_gmail", [])
    par_id = {str(f.get("id")): f for f in filtres}
    if sorted(ordre) != sorted(par_id):
        return jsonify({"ok": False, "error": "Liste d'ids invalide"}), 400
    nouveaux = []
    for i, fid in enumerate(ordre, 1):
        f = par_id[str(fid)]
        f["priority"] = i
        nouveaux.append(f)
    cfg["filtres_gmail"] = nouveaux
    sauver_config(cfg)
    return jsonify({"ok": True, "filtres": nouveaux})


@bp.route("/api/filtres/<int:idx>/toggle", methods=["POST"])
def api_filtres_toggle(idx):
    """Active / met en pause une règle."""
    cfg = charger_config()
    filtres = cfg.get("filtres_gmail", [])
    if not 0 <= idx < len(filtres):
        return jsonify({"ok": False, "error": "Introuvable"}), 404
    filtres[idx]["enabled"] = not filtres[idx].get("enabled", True)
    sauver_config(cfg)
    return jsonify({"ok": True, "filtres": filtres})


# APPRENTISSAGE : l'utilisateur re classe un mail mal rangé
# -----------------------------------------------------------
@bp.route("/api/filtres/corriger", methods=["POST"])
def api_filtres_corriger():
    """L'utilisateur dit : « les mails de X vont dans le libellé Y ».
    L'app crée/met à jour une règle expéditeur → libellé et note
    le nombre de corrections (pour les suggestions et l'alerte)."""
    data = request.get_json() or {}
    email = _txt(data.get("email")).strip().lower()
    label = _txt(data.get("label")).strip()
    if not email or "@" not in email or not label:
        return jsonify({"ok": False, "error": "Email et libellé requis"}), 400

    cfg = charger_config()
    apprentissage = cfg.setdefault("apprentissage", {})
    entree = apprentissage.get(email, {"label": label, "corrections": 0})
    meme_label = entree.get("label") == label
    entree["label"] = label
    entree["corrections"] = 0 if meme_label else entree.get("corrections", 0) + 1
    entree["derniere"] = maintenant_str()
    apprentissage[email] = entree

    # La règle expéditeur devient prioritaire (priority 0 → gagne sur tout)
    filtres = cfg.get("filtres_gmail", [])
    regle = None
    for f in filtres:
        items = f.get("items") or []
        if (f.get("source") == "apprentissage"
                and any(it.get("valeur", "").lower() == email for it in items)):
            regle = f
            break
    if regle is None:
        regle = {"id": f"rule-app-{uuid.uuid4().hex[:6]}", "name": label,
                 "icon": "🧠", "color": "#16a085", "source": "apprentissage"}
        filtres.append(regle)
    regle.update({
        "priority": 0,                    # apprend plus fort que les règles fixes
        "enabled":  True,
        "items":    [{"champ": "from", "valeur": email}],
        "label":    label,
        "name":     label,
    })
    acts = ["label"]
    if _txt(data.get("notify")).lower() in ("1", "true", "oui"):
        acts.append("notify")
    regle["actions"] = acts
    cfg["filtres_gmail"] = sorted(filtres, key=lambda f: f.get("priority", 999))
    sauver_config(cfg)

    suggestions = _calculer_suggestions(cfg)
    return jsonify({"ok": True, "filtres": cfg["filtres_gmail"],
                    "apprentissage": apprentissage, "suggestions": suggestions})


def _calculer_suggestions(cfg):
    """Règles souvent corrigées (>20 % de corrections) → à signaler au front.
    Un expéditeur corrigé 3+ fois en changeant de libellé → suggestion de règle."""
    suggestions = []
    for email, e in cfg.get("apprentissage", {}).items():
        if e.get("corrections", 0) >= 3:
            suggestions.append({"email": email, "label": e.get("label", ""),
                                "corrections": e["corrections"],
                                "type": "expediteur_frequent"})
    for f in cfg.get("filtres_gmail", []):
        total = f.get("nb_classements", 0)
        corr = f.get("nb_corrections", 0)
        if total >= 10 and corr / total > 0.2:
            suggestions.append({"regle": f.get("name", ""), "id": f.get("id"),
                                "taux": round(corr / total, 2),
                                "type": "regle_faible"})
    return suggestions


@bp.route("/api/filtres/suggestions", methods=["GET"])
def api_filtres_suggestions():
    cfg = charger_config()
    return jsonify({"ok": True, "suggestions": _calculer_suggestions(cfg),
                    "apprentissage": cfg.get("apprentissage", {})})


# -----------------------------------------------------------
# NOTIFICATIONS (mails classés Travail / Opportunité…)
# -----------------------------------------------------------
@bp.route("/api/notifications", methods=["GET"])
def api_notifications_get():
    cfg = charger_config()
    notifs = cfg.get("notifications", [])
    vues = session.get("notifs_vues", 0)
    return jsonify({"ok": True, "notifications": notifs[:20],
                    "non_vues": max(0, len(notifs) - vues)})


@bp.route("/api/notifications/vue", methods=["POST"])
def api_notifications_vue():
    cfg = charger_config()
    session["notifs_vues"] = len(cfg.get("notifications", []))
    return jsonify({"ok": True})


@bp.route("/api/filtres/<int:idx>", methods=["DELETE"])
def api_filtres_delete(idx):
    cfg = charger_config()
    filtres = cfg.get("filtres_gmail", [])
    if not 0 <= idx < len(filtres):
        return jsonify({"ok": False, "error": "Introuvable"}), 404
    filtres.pop(idx)
    for i, f in enumerate(filtres, 1):
        f["priority"] = i
    cfg["filtres_gmail"] = filtres
    sauver_config(cfg)
    return jsonify({"ok": True, "filtres": filtres})


@bp.route("/api/filtres/creer-libelles", methods=["POST"])
def api_filtres_creer_libelles():
    """Pré-crée DANS Gmail les libellés de toutes les règles (label action),
    même s'aucun mail ne correspond encore."""
    import imaplib
    cfg = charger_config()
    comptes = get_comptes_actifs()
    if not comptes:
        return jsonify({"ok": False, "error": "Aucun compte actif"}), 400
    libelles = sorted({(f.get("label") or "").strip()
                       for f in cfg.get("filtres_gmail", [])
                       if "label" in (f.get("actions") or []) and (f.get("label") or "").strip()})
    if not libelles:
        return jsonify({"ok": False, "error": "Aucune règle avec libellé"}), 400
    resultats = []
    for compte in comptes:
        mail = None
        crees = []
        try:
            mail = imaplib.IMAP4_SSL("imap.gmail.com", 993, timeout=25)
            mail.login(compte["email"], crypto.dechiffrer(compte["password"]))
            for lbl in libelles:
                # existe déjà ? (le nom IMAP est encodé UTF-7)
                try:
                    status, _ = mail.select(_utf7(lbl))
                    if status == "OK":
                        continue   # existe déjà
                except Exception:
                    pass
                if _creer_libelle(mail, lbl):
                    crees.append(lbl)
            resultats.append({"email": compte["email"], "ok": True, "crees": crees})
            print(f"[LABELS] {compte['email']}: {len(crees)} libellé(s) créé(s)")
        except Exception as e:
            resultats.append({"email": compte["email"], "ok": False, "erreur": str(e)})
        finally:
            if mail:
                try: mail.logout()
                except Exception: pass
    return jsonify({"ok": True, "libelles": libelles, "resultats": resultats})


@bp.route("/api/filtres/appliquer", methods=["POST"])
def api_filtres_appliquer():
    """Exécute maintenant toutes les règles sur toutes les boîtes actives."""
    cfg = charger_config()
    filtres = cfg.get("filtres_gmail", [])
    if not filtres:
        return jsonify({"ok": False, "error": "Aucune règle définie"})
    comptes = get_comptes_actifs()
    if not comptes:
        return jsonify({"ok": False, "error": "Aucun compte actif"})
    uid = current_uid()
    resultats = []
    for compte in comptes:
        resultats.extend(appliquer_filtres_compte(compte, filtres, uid=uid))
    return jsonify({"ok": True, "resultats": resultats})

# ------------------------------------------------------------
# MARQUER RÉPONSE
# ------------------------------------------------------------
@bp.route("/api/mark-response", methods=["POST"])
def api_mark_response():
    data = request.get_json() or {}
    company = (data.get("company") or "").strip()
    repondu = bool(data.get("repondu", True))
    log = charger_log()
    if company not in log:
        return jsonify({"ok": False, "error": "Introuvable"}), 404
    log[company]["repondu"] = repondu
    if repondu: log[company]["date_reponse"] = maintenant_str()
    else: log[company].pop("date_reponse", None)
    sauver_log(log)
    return jsonify({"ok": True, "repondu": repondu})
# ------------------------------------------------------------
# RESET / EXPORTS
# ------------------------------------------------------------
@bp.route("/api/reset-log", methods=["POST"])
def api_reset_log():
    db.log_reset(backup_to=LOG_FILE.with_suffix(".backup.json"),
                 uid=current_uid())
    with _send_lock:
        _send_timestamps[:] = []
        _last_send_time[0] = 0.0
    return jsonify({"ok": True})

@bp.route("/api/export-csv")
def api_export_csv():
    log = charger_log()
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Entreprise", "Email", "Statut", "Réponse", "Bounce",
                "Sans PJ", "Date", "Erreur"])
    for company, info in sorted(log.items(),
                                key=lambda x: x[1].get("date", ""),
                                reverse=True):
        w.writerow([
            company, info.get("email", ""),
            "Envoyé" if info.get("envoye") else "Échec",
            "Oui" if info.get("repondu") else "Non",
            "Oui" if info.get("bounce") else "Non",
            "Oui" if info.get("sans_pj") else "Non",
            info.get("date", ""), info.get("erreur", ""),
        ])
    return Response(out.getvalue().encode("utf-8-sig"),
                    mimetype="text/csv",
                    headers={"Content-Disposition":
                             f'attachment; filename="envois_{datetime.now():%Y%m%d_%H%M}.csv"'})

@bp.route("/api/export-json")
def api_export_json():
    data = {
        "export_le": datetime.now().isoformat(),
        "config":    {k: v for k, v in charger_config().items() if k != "comptes"},
        "log":       charger_log(),
    }
    return Response(
        json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"),
        mimetype="application/json",
        headers={"Content-Disposition":
                 f'attachment; filename="backup_{datetime.now():%Y%m%d_%H%M}.json"'})

@bp.route("/api/import-json", methods=["POST"])
def api_import_json():
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "Aucun fichier"}), 400
    f = request.files["file"]
    try:
        data = json.loads(f.read().decode("utf-8"))
    except Exception as e:
        return jsonify({"ok": False, "error": f"JSON invalide : {e}"}), 400

    imported = data.get("log") or data
    if not isinstance(imported, dict):
        return jsonify({"ok": False, "error": "Structure invalide"}), 400

    current = charger_log()
    avant = len(current)
    current.update(imported)
    sauver_log(current)
    return jsonify({
        "ok": True,
        "importees": len(imported),
        "total_avant": avant,
        "total_apres": len(current),
    })

@bp.route("/api/export-excel")
def api_export_excel():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"ok": False,
                        "error": "openpyxl requis : pip install openpyxl"}), 500

    log = charger_log()
    cfg = charger_config()

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2563EB")
    center = Alignment(horizontal="center", vertical="center")

    ws1 = wb.active
    ws1.title = "Envois"
    headers = ["Entreprise", "Email", "Statut", "Réponse", "Bounce",
               "Sans PJ", "Date envoi", "Date réponse", "Erreur"]
    ws1.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws1.cell(row=1, column=col)
        c.font = header_font; c.fill = header_fill; c.alignment = center

    for company, info in sorted(log.items(),
                                key=lambda x: x[1].get("date", ""),
                                reverse=True):
        ws1.append([
            company, info.get("email", ""),
            "Envoyé" if info.get("envoye") else "Échec",
            "Oui" if info.get("repondu") else "Non",
            "Oui" if info.get("bounce") else "Non",
            "Oui" if info.get("sans_pj") else "Non",
            info.get("date", ""), info.get("date_reponse", ""),
            info.get("erreur", ""),
        ])

    widths = [30, 35, 12, 12, 12, 12, 18, 18, 40]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Statistiques")
    total    = len(log)
    envoyes  = len([x for x in log.values() if x.get("envoye")])
    repondus = len([x for x in log.values() if x.get("repondu")])
    bounces  = len([x for x in log.values() if x.get("bounce")])
    echecs   = total - envoyes

    ws2.append(["Statistique", "Valeur"])
    for col in (1, 2):
        c = ws2.cell(row=1, column=col)
        c.font = header_font; c.fill = header_fill; c.alignment = center

    ws2.append(["Total entreprises contactées", total])
    ws2.append(["Emails envoyés avec succès", envoyes])
    ws2.append(["Échecs d'envoi", echecs])
    ws2.append(["Réponses reçues", repondus])
    ws2.append(["Bounces (emails invalides)", bounces])
    ws2.append(["Taux de réponse (%)",
                round(100 * repondus / envoyes, 1) if envoyes else 0])

    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 15

    ws3 = wb.create_sheet("Candidat")
    ws3.append(["Champ", "Valeur"])
    for col in (1, 2):
        c = ws3.cell(row=1, column=col)
        c.font = header_font; c.fill = header_fill; c.alignment = center

    for k, v in cfg["candidat"].items():
        ws3.append([k.capitalize(), v])
    ws3.append(["Template", cfg["parametres"]["template"]])
    ws3.append(["Délai", f"{get_delai()}s"])
    ws3.append(["Export", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 45

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"candimail_{datetime.now():%Y%m%d_%H%M}.xlsx",
    )
